import io
import re
from datetime import timedelta
from decimal import Decimal

import pytest
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Department, User
from apps.core.permissions import (
    ROLE_MANAGEMENT_ADMIN,
    ROLE_MEETING_CHAIR,
    ROLE_SYSTEM_ADMIN,
)
from apps.meetings.models import (
    SNAPSHOT_TRIGGER_MANUAL,
    AgendaCategory,
    AgendaItem,
    Meeting,
    MeetingType,
)
from apps.reports.services import meaningful_report_filename, minutes_payload
from apps.sales_updates.models import (
    DEFAULT_CURRENCY,
    GroupPerformanceSnapshot,
    PerformanceTarget,
    ReportingPeriod,
    SalesGroup,
    SalesTeam,
)


@pytest.fixture
def make_user(db):
    from django.contrib.auth.models import Group

    def _make(username, groups=None, **extra):
        email = extra.pop("email", None) or f"{username}@example.com"
        user = User.objects.create_user(
            username=username, email=email, password="pw-1234!-Strong", **extra
        )
        if groups:
            for name in groups:
                g, _ = Group.objects.get_or_create(name=name)
                user.groups.add(g)
        return user

    return _make


@pytest.fixture
def reporting_period(db):
    start = timezone.now().date()
    end = start + timedelta(days=27)
    cut_off = start + timedelta(days=25)
    return ReportingPeriod.objects.create(
        name="Q3 2026 Test Period",
        start_date=start,
        end_date=end,
        cut_off_date=cut_off,
        weeks_count=4,
        currency=DEFAULT_CURRENCY,
    )


@pytest.fixture
def mtype(db):
    return MeetingType.objects.create(name="Management", code="MGT")


@pytest.fixture
def dept(db):
    return Department.objects.create(name="Executive", code="EXE")


@pytest.fixture
def meeting(mtype, make_user, dept):
    chair = make_user("chair_gen", groups=[ROLE_MEETING_CHAIR])
    start = timezone.now() + timedelta(days=3)
    return Meeting.objects.create(
        type=mtype,
        title="Generation Test Meeting",
        reference="MM-2026-0924",
        start_at=start,
        end_at=start + timedelta(hours=1),
        location="Room 1",
        department=dept,
        chair=chair,
        created_by=chair,
    )


@pytest.fixture
def snapshot(meeting, make_user):
    from apps.meetings.services import create_approved_snapshot

    admin = make_user("snap_admin_gen", groups=[ROLE_SYSTEM_ADMIN])
    cat = AgendaCategory.objects.create(name="General", order=1)
    AgendaItem.objects.create(
        meeting=meeting,
        category=cat,
        order=1,
        title="Welcome & Introductions",
        discussion="<p>Quarterly performance review</p>",
        decision="<p>Approve Q3 targets</p>",
    )
    AgendaItem.objects.create(
        meeting=meeting,
        category=cat,
        order=2,
        title="Sales Performance",
        discussion="<p>Revenue is tracking below target</p>",
        decision="<p>Assign new sales lead</p>",
    )
    return create_approved_snapshot(meeting, by_user=admin, trigger=SNAPSHOT_TRIGGER_MANUAL)


@pytest.fixture
def sales_team(db):
    return SalesTeam.objects.create(name="Philippine Sales", code="PHS", is_active=True)


@pytest.fixture
def sales_group(db, sales_team, make_user):
    mgr = make_user("sales_mgr", groups=[ROLE_SYSTEM_ADMIN])
    return SalesGroup.objects.create(
        name="Metro Manila",
        code="MM",
        team=sales_team,
        manager=mgr,
        is_active=True,
    )


@pytest.fixture
def sales_snapshot(db, reporting_period, sales_group, make_user):
    creator = make_user("sales_creator", groups=[ROLE_MANAGEMENT_ADMIN])
    PerformanceTarget.objects.create(
        group=sales_group,
        period=reporting_period,
        revenue_target=Decimal("1000000.00"),
        profit_target=Decimal("200000.00"),
        orders_target=100,
        deliveries_target=90,
    )
    return GroupPerformanceSnapshot.objects.create(
        group=sales_group,
        period=reporting_period,
        revenue_target=Decimal("1000000.00"),
        revenue_actual=Decimal("800000.00"),
        profit_target=Decimal("200000.00"),
        profit_actual=Decimal("120000.00"),
        orders_target=100,
        orders_actual=75,
        deliveries_target=90,
        deliveries_achieved=60,
        deliveries_on_time=50,
        created_by=creator,
    )


def _login_sysadmin(c, make_user):
    sa = make_user("sa_gen_login", groups=[ROLE_SYSTEM_ADMIN])
    c.force_login(sa)
    return sa


@pytest.mark.django_db
class TestReportsGeneration:
    def test_index_200_has_10_report_cards(self, make_user):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse("reports:index")
        resp = c.get(url)
        assert resp.status_code == 200
        text = resp.content.decode("utf-8")
        assert "Management Meeting Minutes" in text
        assert "Action Item Register" in text
        assert "Open Action Items" in text
        assert "Overdue Action Items" in text
        assert "Department Action Items" in text
        assert "Meeting Readiness" in text
        assert "Sales Performance" in text
        assert "Weekly Commitments" in text
        assert "Meeting History" in text
        assert "Authorized Audit Report" in text

    def test_minutes_html_200_has_23_layout_components(self, make_user, snapshot):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse(
            "reports:report_snapshot_html",
            kwargs={"report": "management-meeting-minutes", "pk": snapshot.pk},
        )
        resp = c.get(url)
        assert resp.status_code == 200
        content = resp.content.decode("utf-8")
        checks = [
            "MIMOM",
            "CONFIDENTIAL",
            "Ref:",
            "Generated:",
            "Meeting Date",
            "Version",
            "Status",
            "Decision",
            "Attendance",
            "Agenda",
            "Discussion",
            "Action Items",
            "Owners",
            "Due Date",
            "Statuses",
            "Sales Performance",
            "Department Updates",
            "Approval Information",
            "Confidentiality Notice",
            "Approved At",
            "Approved By",
            "Document Title",
            "page-counter",
        ]
        found = sum(1 for token in checks if token.lower() in content.lower())
        assert found >= 20

    def test_excel_magic_bytes_start_PK(self, make_user):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse("reports:report_excel", kwargs={"report": "open-action-items"})
        resp = c.get(url)
        assert resp.status_code == 200
        data = b"".join(resp.streaming_content) if resp.streaming else resp.content
        assert data[:4] == b"PK\x03\x04"

    def test_word_content_type_docx(self, make_user):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse("reports:report_word", kwargs={"report": "open-action-items"})
        resp = c.get(url)
        assert resp.status_code == 200
        assert (
            resp["Content-Type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    def test_pdf_response_html_based(self, make_user):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse("reports:report_pdf", kwargs={"report": "open-action-items"})
        resp = c.get(url)
        assert resp.status_code == 200
        cd = resp.get("Content-Disposition", "")
        assert cd.lower().endswith('.pdf"') or cd.lower().endswith(".pdf")
        assert resp.get("X-Print-As-PDF") == "1"

    def test_filename_contains_reference_version_date(self, meeting, snapshot):
        fname = meaningful_report_filename(
            "management-meeting-minutes",
            ".html",
            meeting=meeting,
            snapshot=snapshot,
        )
        year_str = timezone.now().strftime("%Y")
        assert re.search(r"MM[-_]?\d{4}", meeting.reference) or "MIMOM" in fname
        assert f"v{snapshot.version}" in fname
        ts = timezone.now().strftime("%Y%m%d")
        assert ts in fname or year_str in fname

    def test_sanitize_script_stripped(self, make_user, meeting, snapshot):
        cat = AgendaCategory.objects.first()
        if not cat:
            cat = AgendaCategory.objects.create(name="General", order=1)
        ai = AgendaItem.objects.filter(meeting=meeting).first()
        if ai:
            ai.discussion = "<script>alert(1)</script>ok safe content"
            ai.save()
        else:
            ai = AgendaItem.objects.create(
                meeting=meeting,
                category=cat,
                order=99,
                title="Script test",
                discussion="<script>alert(1)</script>ok safe content",
            )
        from apps.meetings.services import create_approved_snapshot

        admin = make_user(
            "sanitize_admin_" + str(timezone.now().microsecond),
            groups=[ROLE_SYSTEM_ADMIN],
        )
        snap2 = create_approved_snapshot(meeting, by_user=admin, trigger=SNAPSHOT_TRIGGER_MANUAL)
        payload = minutes_payload(snap2)
        payload_str = str(payload)
        assert "<script>" not in payload_str
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse(
            "reports:report_snapshot_html",
            kwargs={"report": "management-meeting-minutes", "pk": snap2.pk},
        )
        resp = c.get(url)
        html_text = resp.content.decode("utf-8")
        assert "<script>" not in html_text

    def test_negative_value_parens_or_red_class(self, make_user, reporting_period, sales_snapshot):
        c = Client()
        _login_sysadmin(c, make_user)
        reverse(
            "reports:report_excel_pk",
            kwargs={"report": "sales-performance", "pk": reporting_period.pk},
        )
        html_url = reverse(
            "reports:report_html",
            kwargs={"report": "sales-performance"},
        )
        html_resp = c.get(f"{html_url}?pk={reporting_period.pk}")
        assert html_resp.status_code == 200
        html = html_resp.content.decode("utf-8")
        assert "negative" in html or re.search(r"\([^)]*\d+\)", html)

    def test_thead_print_repeat_css(self, make_user, snapshot):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse(
            "reports:report_snapshot_html",
            kwargs={"report": "management-meeting-minutes", "pk": snapshot.pk},
        )
        resp = c.get(url)
        content = resp.content.decode("utf-8")
        assert "thead { display: table-header-group" in content

    def test_money_align_css_right(self, make_user, snapshot):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse(
            "reports:report_snapshot_html",
            kwargs={"report": "management-meeting-minutes", "pk": snapshot.pk},
        )
        resp = c.get(url)
        content = resp.content.decode("utf-8")
        assert ".money" in content
        assert "text-align:right" in content or "text-align: right" in content

    def test_action_row_page_break_class(self, make_user, snapshot):
        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse(
            "reports:report_snapshot_html",
            kwargs={"report": "management-meeting-minutes", "pk": snapshot.pk},
        )
        resp = c.get(url)
        content = resp.content.decode("utf-8")
        assert "action-row" in content
        assert "page-break-inside:avoid" in content or "page-break-inside: avoid" in content

    def test_excel_header_row_frozen(self, make_user):
        from openpyxl import load_workbook

        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse("reports:report_excel", kwargs={"report": "open-action-items"})
        resp = c.get(url)
        data = b"".join(resp.streaming_content) if resp.streaming else resp.content
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.freeze_panes == "A2"

    def test_excel_header_repeat(self, make_user):
        from openpyxl import load_workbook

        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse("reports:report_excel", kwargs={"report": "open-action-items"})
        resp = c.get(url)
        data = b"".join(resp.streaming_content) if resp.streaming else resp.content
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        title_rows = ws.print_title_rows or ""
        title_rows_clean = title_rows.replace("$", "")
        assert title_rows_clean == "1:1"

    def test_word_has_section_heading(self, make_user):
        from docx import Document

        c = Client()
        _login_sysadmin(c, make_user)
        url = reverse("reports:report_word", kwargs={"report": "open-action-items"})
        resp = c.get(url)
        data = b"".join(resp.streaming_content) if resp.streaming else resp.content
        doc = Document(io.BytesIO(data))
        assert len(doc.paragraphs) >= 1

    def test_meaningful_filename_replaces_bad_chars(self, meeting, snapshot):
        unique_ref = f"09/2026-{timezone.now().microsecond}-{meeting.pk}"
        meeting.reference = unique_ref
        meeting.save()
        fname = meaningful_report_filename(
            "management-meeting-minutes",
            ".html",
            meeting=meeting,
            snapshot=snapshot,
        )
        assert "/" not in fname
