from __future__ import annotations

import io
from collections.abc import Iterable

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - openpyxl is required at runtime
    Workbook = None  # type: ignore[assignment]
    Font = None  # type: ignore[assignment]
    PatternFill = None  # type: ignore[assignment]
    Alignment = None  # type: ignore[assignment]
    get_column_letter = None  # type: ignore[assignment]


_HEADER_FONT_KWARGS = dict(bold=True, color="FFFFFF", name="Calibri", size=11)
_HEADER_FILL_KWARGS = dict(start_color="FF1F3A68", end_color="FF1F3A68", fill_type="solid")
_MONEY_FORMAT = '_-"₱"* #,##0.00_-;[Red]-"₱"* #,##0.00_-'
_MONEY_ALIGN_KWARGS = dict(horizontal="right", vertical="top")


def write_tabular_report(
    ws,
    title: str,
    headers: list[str],
    rows: list[list],
    money_cols: Iterable[int] | frozenset[int] = frozenset(),
):
    if Workbook is None:  # pragma: no cover
        raise RuntimeError("openpyxl is required for Excel report generation")

    money = frozenset(money_cols or ())

    ws.title = (title or "Sheet1")[:31]

    header_font = Font(**_HEADER_FONT_KWARGS)
    header_fill = PatternFill(**_HEADER_FILL_KWARGS)
    money_align = Alignment(**_MONEY_ALIGN_KWARGS)

    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for row in rows:
        ws.append(list(row))

    if money:
        for col_idx in sorted(money):
            col_letter = get_column_letter(col_idx + 1)
            for row_idx in range(2, ws.max_row + 1):
                cell = ws[f"{col_letter}{row_idx}"]
                cell.number_format = _MONEY_FORMAT
                cell.alignment = money_align

    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"

    widths: dict[int, int] = {}
    for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for idx, cell in enumerate(row_cells):
            value = "" if cell.value is None else str(cell.value)
            length = len(value)
            if idx not in widths or length > widths[idx]:
                widths[idx] = min(length + 2, 60)

    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx + 1)].width = max(width, 10)

    ws.sheet_properties.pageSetUpPr = None
    return ws


def write_workbook_bytes(
    sheets: dict[str, tuple[list[str], list[list], Iterable[int] | None]],
) -> bytes:
    if Workbook is None:  # pragma: no cover
        raise RuntimeError("openpyxl is required for Excel report generation")

    wb = Workbook()
    first = True
    for sheet_name, payload in sheets.items():
        headers, rows, money_cols = payload
        money = frozenset(money_cols or ())
        if first:
            ws = wb.active
            first = False
        else:
            ws = wb.create_sheet()
        write_tabular_report(ws, sheet_name, list(headers), list(rows), money)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
