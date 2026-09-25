from __future__ import annotations

import io

try:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:  # pragma: no cover - python-docx is required at runtime
    Document = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    WD_ALIGN_VERTICAL = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]


def _bold_header_row(row, cols: int):
    for c in range(cols):
        cell = row.cells[c]
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
            if not paragraph.runs:
                run = paragraph.add_run()
                run.bold = True


def _apply_money_alignment(cell):
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def build_docx_report(
    title: str,
    subtitle: str | None = None,
    sections: list[dict] | list[tuple] | None = None,
) -> bytes:
    if Document is None:  # pragma: no cover
        raise RuntimeError("python-docx is required for Word report generation")

    doc = Document()

    styles = doc.styles
    normal_style = styles["Normal"]
    normal_font = normal_style.font
    normal_font.name = "Georgia"
    normal_font.size = Pt(10)

    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.name = "Georgia"

    if subtitle:
        sub_para = doc.add_paragraph(subtitle)
        for run in sub_para.runs:
            run.italic = True

    if not sections:
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    normalized_sections: list[dict] = []
    for section in sections:
        if isinstance(section, dict):
            normalized_sections.append(section)
        elif isinstance(section, (tuple, list)):
            if len(section) >= 2:
                normalized_sections.append(
                    {
                        "heading": section[0],
                        "rows": section[1],
                        "money_cols": section[2] if len(section) > 2 else frozenset(),
                    }
                )
            else:
                normalized_sections.append(
                    {"heading": str(section[0]) if section else "", "rows": []}
                )

    for sec in normalized_sections:
        heading = sec.get("heading", "")
        rows = sec.get("rows") or []
        money_cols = frozenset(sec.get("money_cols") or ())

        if heading:
            doc.add_heading(str(heading), level=2)

        if not rows or (len(rows) == 1 and all(v is None or str(v) == "" for v in rows[0])):
            continue

        n_cols = len(rows[0]) if rows else 0
        if n_cols == 0:
            continue

        n_rows = len(rows)
        table = doc.add_table(rows=n_rows, cols=n_cols)
        table.style = "Table Grid"
        table.autofit = True

        for r_idx in range(n_rows):
            row = rows[r_idx]
            tr = table.rows[r_idx]
            for c_idx in range(n_cols):
                cell = tr.cells[c_idx]
                val = row[c_idx] if c_idx < len(row) else ""
                cell.text = "" if val is None else str(val)

                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.font.name = "Georgia"

                if r_idx == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                        if not paragraph.runs:
                            run = paragraph.add_run()
                            run.bold = True

                if c_idx in money_cols and r_idx > 0:
                    _apply_money_alignment(cell)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
